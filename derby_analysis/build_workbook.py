import json, re, sys, hashlib
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.formatting.rule import CellIsRule

HERE = Path(__file__).resolve().parent
IN = HERE / "inputs"
OUT = sys.argv[1] if len(sys.argv) > 1 else str(HERE.parent / "Derby_Full_EGD_PO_Brand_Preference_Content_Analysis.xlsx")
data = json.load(open(IN / "analysis_content.json", encoding="utf-8"))
coverage = json.load(open(IN / "egd_section_coverage.json", encoding="utf-8"))
prior = json.load(open(IN / "prior_analysis_extract.json", encoding="utf-8"))
audit = json.load(open(IN / "transcript_audit.json", encoding="utf-8"))
names = sorted(f["name"] for f in audit)

def find(p):
    m = [n for n in names if re.match(p, n, re.I)]
    assert len(m) == 1, (p, m)
    return m[0]

cur = dict(DM3=find(r"^Checked_Dm3_EGD_"), DM6=find(r"^Checked_DM6_EGDs_"), RM5=find(r"^Checked_RM5_EGD_"),
           RM2=find(r"^RM2_EGD_"), POD3=find(r"^Checked_PO-D3_"), R5PO=find(r"^Checked_R5_PO_"),
           R2PO=find(r"^R2_PO_"), POD06=find(r"^PO_D06_"))

evidence = [r[:7] + [cur[r[0]]] for r in data["contextEvidence"] + prior["evidence"]]
po_src = {"PO-D3": cur["POD3"], "R5": cur["R5PO"], "R2-PO": cur["R2PO"]}
po_rows = [r + ["Cigarette Brand Preference module only", po_src.get(r[0], cur["POD06"])] for r in prior["po"]]

def family(s):
    s = s.lower()
    if re.search(r"intro|daily|dream|perfect|lifestyle|phych|challenge|entertain|media|shopping|spending|collecting|celebrity", s): return "Life, aspiration and media"
    if re.search(r"category|smoker|consum|word game|purchase$", s): return "Category and consumption"
    if re.search(r"brand|alternative|reject|selection|unique|persona|pack|stick|strength|weak|price|quality|appeal|stimulus|name", s): return "Brand and product"
    return "Awareness, activation and future"

section_cov = [[r["group"], family(r["section"]), r["section"], r["location"],
                r["records"] if r["records"] is not None else "n.a.", r["included"], r["source"]] for r in coverage]

C = dict(navy="17324D", red="C6342E", teal="2A7F8E", green="548235", paleBlue="EAF1F8", paleGold="FFF2CC",
         paleRed="FCE4D6", paleGreen="E2F0D9", grey="F2F2F2", white="FFFFFF", text="222222")
F = "Arial"
fill = lambda c: PatternFill("solid", fgColor=c)
thin_w = Side(style="thin", color="FFFFFF"); thin_g = Side(style="thin", color="D9D9D9")
box = lambda s: Border(left=s, right=s, top=s, bottom=s)

def style(ws, rng, font=None, fl=None, align=None, border=None, height=None, numfmt=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fl: c.fill = fl
            if align: c.alignment = align
            if border: c.border = border
            if numfmt: c.number_format = numfmt
        if height: ws.row_dimensions[row[0].row].height = height

def setup(ws, title, subtitle, ncols):
    ws.sheet_view.showGridLines = False
    ws["A2"] = title; ws["A2"].font = Font(name=F, size=16, bold=True, color=C["navy"])
    for i in range(1, ncols + 1):
        ws.cell(2, i).border = Border(bottom=Side(style="medium", color=C["navy"]))
    ws["A3"] = subtitle; ws["A3"].font = Font(name=F, size=9, italic=True, color="566573")
    ws.row_dimensions[3].height = 30

def table(ws, start, headers, rows, widths, body_h=46, hfill=C["navy"]):
    n = len(headers); end = start + len(rows); ec = col(n)
    for j, h in enumerate(headers, 1): ws.cell(start, j, h)
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r, 1): ws.cell(start + i, j, v)
    style(ws, f"A{start}:{ec}{end}", font=Font(name=F, size=9, color=C["text"]),
          align=Alignment(vertical="top", wrap_text=True), border=box(thin_g))
    style(ws, f"A{start}:{ec}{start}", font=Font(name=F, size=10, bold=True, color=C["white"]), fl=fill(hfill),
          align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=box(thin_w), height=34)
    for r in range(start + 1, end + 1): ws.row_dimensions[r].height = body_h
    for i, w in enumerate(widths, 1): ws.column_dimensions[col(i)].width = w
    ws.freeze_panes = f"A{start + 1}"
    ws.auto_filter.ref = f"A{start}:{ec}{end}"
    return end

wb = Workbook()

# Summary
s = wb.active; s.title = "Summary"
setup(s, "Derby full EGD content analysis", "Full EGD discussions are primary evidence. PO evidence is restricted to the Cigarette Brand Preference module.", 8)
hdr = ["Evidence base", "Value proposition", "Main vulnerability", "Conversion condition", "Status tension", "Full EGD coverage", "PO boundary", "Method"]
val = ["4 complete EGDs\n~28 participants", "Affordable, strong and smoky everyday satisfaction", "Persistent odor, then filter and burn execution",
       "Same price, Derby-like feel, better filter and less odor", "Everyday fit is stronger than visible-status fit",
       f"{len(section_cov)} source sections or paragraph bands included", "4 PO interviews; brand preference section only", "Qualitative group-level thematic analysis"]
for j, (h, v) in enumerate(zip(hdr, val), 1): s.cell(5, j, h); s.cell(6, j, v)
style(s, "A5:H5", font=Font(name=F, size=10, bold=True, color=C["white"]), fl=fill(C["navy"]),
      align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=box(thin_w), height=30)
style(s, "A6:H6", font=Font(name=F, size=10, bold=True, color=C["text"]), fl=fill(C["paleBlue"]),
      align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=box(thin_g), height=78)
main = [("1", "Life context", "Participants are managing study, work, family expectation and limited cash while trying to progress socially and economically."),
        ("2", "Role of smoking", "Smoking is attached to morning/night routines, tea, meals, work breaks, stress relief, focus, waiting and adda."),
        ("3", "Why Derby wins", "Derby combines frequent-use affordability with strong regular delivery, visible smoke, availability and learned physical fit."),
        ("4", "Why Derby is vulnerable", "Odor reduces social acceptability. Filter softness/heat, paper, burn and ash issues signal lower execution quality."),
        ("5", "Competitive risk", "At price parity, Star and premium brands gain. Some smokers already trade up for guests, campus or visible occasions."),
        ("6", "How trial converts", "Peers and retailers are more credible than poster copy. Product must retain Derby-like satisfaction while improving filter and odor.")]
s["A9"], s["B9"], s["C9"] = "No.", "Area", "Finding"; s.merge_cells("C9:H9")
style(s, "A9:H9", font=Font(name=F, size=10, bold=True, color=C["white"]), fl=fill(C["red"]), border=box(thin_w))
for i, (a, b, c) in enumerate(main):
    r = 10 + i; s[f"A{r}"], s[f"B{r}"], s[f"C{r}"] = a, b, c; s.merge_cells(f"C{r}:H{r}")
style(s, "A10:H15", font=Font(name=F, size=10, color=C["text"]), align=Alignment(vertical="center", wrap_text=True), border=box(thin_g), height=46)
style(s, "A10:A15", font=Font(name=F, size=11, bold=True, color=C["red"]), fl=fill(C["paleRed"]), align=Alignment(horizontal="center", vertical="center"))
s["A18"], s["B18"], s["F18"] = "Priority", "Recommended action", "Reason"; s.merge_cells("B18:E18"); s.merge_cells("F18:H18")
style(s, "A18:H18", font=Font(name=F, size=10, bold=True, color=C["white"]), fl=fill(C["teal"]), border=box(thin_w))
actions = [("P1", "Reduce lingering odor without reducing strength, smoke volume or taste.", "Odor is the most consistent cross-group product and social barrier."),
           ("P2", "Improve filter-end integrity and burn/paper/ash consistency.", "These are repeated, tangible quality signals."),
           ("P3", "Defend relative value against Star and premium brands before price convergence.", "Equal pricing weakens practical loyalty."),
           ("P4", "Use peer/retailer-led product trial around a clear proof.", "Recommendation and experience are more credible than passive claims.")]
for i, (a, b, c) in enumerate(actions):
    r = 19 + i; s[f"A{r}"], s[f"B{r}"], s[f"F{r}"] = a, b, c; s.merge_cells(f"B{r}:E{r}"); s.merge_cells(f"F{r}:H{r}")
style(s, "A19:H22", font=Font(name=F, size=10, color=C["text"]), align=Alignment(vertical="center", wrap_text=True), border=box(thin_g), height=52)
style(s, "A19:A22", font=Font(name=F, size=10, bold=True, color=C["green"]), fl=fill(C["paleGreen"]), align=Alignment(horizontal="center", vertical="center"))
for i, w in enumerate([10, 18, 18, 18, 18, 18, 18, 18], 1): s.column_dimensions[col(i)].width = w

# Themes
s = wb.create_sheet("Full EGD Themes")
setup(s, "Themes across the complete EGDs", "Context, aspiration, media, shopping, category, consumption, brand, product and activation sections are included.", 10)
rows = [[t[0], t[1], t[3], *t[2], None, None, t[4]] for t in data["themes"]]
end = table(s, 5, ["Code", "Theme", "Evidence summary", "DM3", "DM6", "RM5", "RM2", "Groups", "% groups", "Implication"], rows,
            [9, 34, 58, 8, 8, 8, 8, 10, 11, 58], 60)
for r in range(6, end + 1):
    s[f"H{r}"] = f"=SUM(D{r}:G{r})"; s[f"I{r}"] = f"=H{r}/4"
style(s, f"D6:I{end}", align=Alignment(horizontal="center", vertical="top", wrap_text=True))
style(s, f"I6:I{end}", numfmt="0%")
s.conditional_formatting.add(f"D6:G{end}", CellIsRule(operator="equal", formula=["1"], fill=fill(C["paleGreen"]), font=Font(bold=True, color=C["green"])))
s.conditional_formatting.add(f"D6:G{end}", CellIsRule(operator="equal", formula=["0"], fill=fill(C["grey"]), font=Font(color="777777")))
s[f"A{end + 2}"] = "Counting rule: 1 = theme present in that EGD group, 0 = not evidenced. Prevalence is per group, not per participant."
s[f"A{end + 2}"].font = Font(name=F, size=9, italic=True, color="566573")

# Group profiles
s = wb.create_sheet("EGD Group Profiles")
setup(s, "EGD group profiles", "Each row summarizes the complete discussion before interpreting Derby preference.", 12)
table(s, 5, ["Group", "City", "Segment", "Age", "SEC", "Participants", "Life context", "Dreams and pressures", "Media and culture",
             "Shopping and money", "Smoking and Derby role", "Main tension"], data["groupProfiles"], [9, 15, 14, 10, 10, 11, 36, 38, 38, 38, 40, 40], 112)

# Section coverage
s = wb.create_sheet("EGD Section Coverage")
setup(s, "Full EGD coverage register", "All tagged sections from the three checked EGD spreadsheets and all paragraph bands from the RM2 EGD DOCX are included.", 7)
table(s, 5, ["Group", "Section family", "Source section", "Source location", "Transcript records", "Included", "Current source file"],
      section_cov, [9, 28, 34, 20, 16, 11, 70], 34)

# Evidence
s = wb.create_sheet("EGD Evidence")
setup(s, "EGD evidence log", "Participant evidence across the full EGD scope. Moderator questions are not treated as findings.", 8)
table(s, 5, ["Group", "Theme", "Speaker", "Source row or paragraph", "Participant verbatim (Bangla)", "Concise English translation",
             "Analytic meaning", "Current source file"], evidence, [9, 19, 12, 20, 54, 50, 50, 70], 72)

# PO
s = wb.create_sheet("PO Brand Preference")
setup(s, "PO comparison: Cigarette Brand Preference only", "Lifestyle, psychography, media and unrelated PO sections are excluded from this sheet and from PO conclusions.", 8)
table(s, 5, ["PO ID", "Profile", "Brand journey and entry", "Derby equity", "Price, fallback and trial", "How it supports the EGD conclusion",
             "PO scope used", "Current source file"], po_rows, [12, 29, 45, 43, 45, 45, 31, 70], 96, C["teal"])

# Implications
s = wb.create_sheet("Implications")
setup(s, "Brand implications", "Actions integrate full-EGD context with the restricted PO brand-preference comparison.", 6)
imp = [[1, "What must Derby retain?", "Affordable strong regular satisfaction, smoke volume, familiarity and easy availability.", "All scoped PO cases reinforce value plus satisfactory delivery.", "Set non-negotiable sensory guardrails before reformulation.", "A lighter or less recognizable Derby could lose its core users."],
       [2, "What should be fixed first?", "Odor, then filter-end integrity and burn/paper/ash consistency.", "POs also use filter, paper and burn as quality cues.", "Prototype and blind-test low-odor, stronger-filter executions.", "Social rejection and visible quality defects sustain trade-up."],
       [3, "How strong is loyalty?", "Habit is strong, but price parity and visible occasions can unlock trade-up.", "PO price tolerance varies, confirming conditional loyalty.", "Model relative price gaps and test parity with Star/premium choices.", "A small architecture change may cause disproportionate switching."],
       [4, "How should trial work?", "Peer/retailer proof, Derby-like feel, better filter and less odor.", "POs also prefer recommendation and one- or two-stick trial.", "Use guided sampling and repeat exposure, not awareness alone.", "Attractive pack without product superiority will not retain."],
       [5, "How can status improve?", "Reduce odor and strengthen quality cues while remaining accessible.", "POs show pack cues help only when product fit is maintained.", "Lift social acceptability without over-premiumizing price or tone.", "Luxury cues alone could alienate the everyday core."]]
table(s, 5, ["Priority", "Business question", "Full EGD answer", "PO triangulation", "Recommended action", "Risk if ignored"], imp, [10, 31, 51, 48, 51, 48], 86, C["red"])

# Source audit (file bytes are not available here; rows/characters/text hash come from transcript_audit.json)
s = wb.create_sheet("Source Audit")
setup(s, "Source audit and scope", "Text rows, characters and a SHA-256 of the extracted transcript text are recorded so later edits or replacement can be detected.", 8)
src = []
for f in sorted(audit, key=lambda f: f["name"]):
    egd = "EGD" in f["name"].upper()
    h = hashlib.sha256("\n".join(r["text"] for r in f["content"]).encode("utf-8")).hexdigest().upper()
    src.append(["Primary" if egd else "Supporting", "EGD" if egd else "PO", f["name"], f["rows"], f["chars"], h,
                "Complete transcript" if egd else "Cigarette Brand Preference section only",
                "Full EGD synthesis and evidence" if egd else "PO triangulation only"])
end = table(s, 5, ["Evidence tier", "Type", "Current filename", "Text rows", "Characters", "SHA-256 of extracted text", "Scope used", "Analytical role"],
            src, [16, 12, 76, 12, 14, 68, 34, 32], 52)
style(s, f"D6:E{end}", numfmt="#,##0")
for k, (lab, txt, fc) in enumerate([("Method note", "The analysis uses deductive categories from the discussion guides and inductive themes repeated across transcripts.", C["paleGold"]),
                                    ("Counting rule", "Theme prevalence is counted once per EGD group. It is not a participant percentage or population estimate.", C["paleBlue"])]):
    r = end + 3 + k; s[f"A{r}"], s[f"B{r}"] = lab, txt; s.merge_cells(f"B{r}:H{r}")
    style(s, f"A{r}:H{r}", font=Font(name=F, size=10, italic=True, color=C["text"]), fl=fill(fc),
          align=Alignment(vertical="center", wrap_text=True), border=box(thin_g), height=44)

wb.save(OUT)
print(json.dumps(dict(coverage=len(section_cov), evidence=len(evidence), themes=len(data["themes"]), po=len(po_rows), sources=len(src))))
